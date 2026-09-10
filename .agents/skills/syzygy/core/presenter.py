"""
SYZYGY Presentation Engine: PPT Master Native Vector Generator
Generates native vector PowerPoint decks (.pptx) adhering to SIH Grand Finale standard.
"""

from typing import Dict, Any, List

class PPTMasterDeckGenerator:
    def __init__(self, title: str = "SYZYGY Presentation"):
        self.title = title
        self.slides: List[Dict[str, Any]] = []

    def add_slide(self, title: str, category: str, bullet_points: List[str], metric_badge: str = ""):
        self.slides.append({
            "slideNumber": len(self.slides) + 1,
            "title": title,
            "category": category,
            "points": bullet_points,
            "badge": metric_badge
        })

    def render_pptx(self, output_path: str = "presentation.pptx") -> str:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor

            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank_layout = prs.slide_layouts[6]

            for s in self.slides:
                slide = prs.slides.add_slide(blank_layout)

                # Background dark canvas
                bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
                bg.fill.solid()
                bg.fill.fore_color.rgb = RGBColor(9, 13, 22) # #090D16
                bg.line.fill.background()

                # Category Supertitle
                catBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
                ctf = catBox.text_frame
                cp = ctf.paragraphs[0]
                cp.text = s["category"].upper()
                cp.font.size = Pt(12)
                cp.font.bold = True
                cp.font.color.rgb = RGBColor(99, 102, 241) # Indigo

                # Main Slide Title
                titleBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
                ttf = titleBox.text_frame
                tp = ttf.paragraphs[0]
                tp.text = s["title"]
                tp.font.size = Pt(32)
                tp.font.bold = True
                tp.font.color.rgb = RGBColor(248, 250, 252)

                # Card Surface
                card = slide.shapes.add_shape(1, Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
                card.fill.solid()
                card.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
                card.line.color.rgb = RGBColor(30, 41, 59)

                card_tf = card.text_frame
                for pt in s["points"]:
                    p = card_tf.add_paragraph()
                    p.text = f"  •  {pt}"
                    p.font.size = Pt(18)
                    p.font.color.rgb = RGBColor(226, 232, 240)

            prs.save(output_path)
            return f"SUCCESS: Generated native vector presentation at '{output_path}'"
        except ImportError:
            import json
            json_path = output_path.replace(".pptx", ".json")
            with open(json_path, "w", encoding="utf-8") as f:
                json.dump({"title": self.title, "slides": self.slides}, f, indent=2)
            return f"python-pptx not available. Outputted slide manifest to '{json_path}'"
